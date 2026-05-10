import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FILES = [
    r"C:\Users\shree\Desktop\ssibm\project_report_part1.md",
    r"C:\Users\shree\Desktop\ssibm\project_report_part2.md",
    r"C:\Users\shree\Desktop\ssibm\project_report_part3.md",
]
OUTPUT = r"C:\Users\shree\Desktop\ssibm\Student_Information_System_Report.docx"

doc = Document()

# ── Page margins ──────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin   = Cm(3.0)
    section.right_margin  = Cm(2.54)

# ── Helper: set font ──────────────────────────────────────────
def set_font(run, name="Times New Roman", size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_page_break(doc):
    doc.add_page_break()

def style_paragraph(para, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                     space_before=0, space_after=6):
    para.alignment = alignment
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after  = Pt(space_after)

# ── Title page ────────────────────────────────────────────────
def add_title_page():
    doc.add_paragraph()
    doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("STUDENT INFORMATION SYSTEM\nWITH AI-BASED ADMISSION ENQUIRY CHATBOT")
    r.font.name = "Times New Roman"
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    doc.add_paragraph()
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = s.add_run("A Project Report\n\nSubmitted in partial fulfillment of the requirements\nfor the award of the degree of\n\nBachelor of Computer Applications (BCA)")
    sr.font.name = "Times New Roman"
    sr.font.size = Pt(12)

    doc.add_paragraph()
    i = doc.add_paragraph()
    i.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ir = i.add_run("Sri Siddhartha Institute of Business Management (SSIBM)\nSSIT Campus, Maralur, Tumakuru, Karnataka – 572105\nAffiliated to Tumkur University\n\nAcademic Year: 2025–2026")
    ir.font.name = "Times New Roman"
    ir.font.size = Pt(12)
    ir.font.bold = True

    doc.add_paragraph()
    g = doc.add_paragraph()
    g.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gr = g.add_run("Submitted By:\n[Student Name]\nRegister Number: [XXXXXXX]\nVI Semester, BCA\n\nGuide:\n[Faculty Name]\nDepartment of Computer Applications")
    gr.font.name = "Times New Roman"
    gr.font.size = Pt(12)
    add_page_break(doc)

add_title_page()

# ── Parse and add markdown content ───────────────────────────
def parse_and_add(filepath):
    with open(filepath, encoding="utf-8") as f:
        lines = f.readlines()

    in_code = False
    code_lines = []

    for line in lines:
        stripped = line.rstrip("\n")

        # Code block toggle
        if stripped.startswith("```"):
            if in_code:
                # End code block — add accumulated lines
                if code_lines:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(1)
                    p.paragraph_format.space_after = Pt(6)
                    run = p.add_run("\n".join(code_lines))
                    run.font.name = "Courier New"
                    run.font.size = Pt(9)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(stripped)
            continue

        # Skip horizontal rules
        if re.match(r"^-{3,}$", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            continue

        # H1
        if stripped.startswith("# ") and not stripped.startswith("## "):
            text = stripped[2:].strip()
            if text and text not in ("PART 1: CHAPTERS 1, 2, 3",
                                     "PART 2: CHAPTERS 4-6",
                                     "PART 3: CHAPTERS 7–10"):
                p = doc.add_heading(text, level=1)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(16)
                    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            continue

        # H2
        if stripped.startswith("## "):
            text = stripped[3:].strip()
            p = doc.add_heading(text, level=2)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(14)
                run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            continue

        # H3
        if stripped.startswith("### "):
            text = stripped[4:].strip()
            p = doc.add_heading(text, level=3)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
                run.font.bold = True
            continue

        # H4
        if stripped.startswith("#### "):
            text = stripped[5:].strip()
            p = doc.add_paragraph(text)
            p.paragraph_format.space_before = Pt(4)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
                run.font.bold = True
            continue

        # Table row
        if stripped.startswith("|"):
            cells = [c.strip() for c in stripped.split("|") if c.strip()]
            if all(re.match(r"^[-:]+$", c) for c in cells):
                continue  # separator row
            # Create or add to table — simple: just add as spaced paragraph
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run("  |  ".join(cells))
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            continue

        # Blockquote
        if stripped.startswith("> "):
            text = stripped[2:].strip()
            p = doc.add_paragraph(text)
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.space_after = Pt(4)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)
                run.font.italic = True
            continue

        # Bullet list
        if re.match(r"^[-*] ", stripped):
            text = stripped[2:].strip()
            # Remove bold markers
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
            p = doc.add_paragraph(text, style="List Bullet")
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)
            continue

        # Numbered list
        if re.match(r"^\d+\. ", stripped):
            text = re.sub(r"^\d+\. ", "", stripped).strip()
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
            p = doc.add_paragraph(text, style="List Number")
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)
            continue

        # Empty line
        if not stripped:
            doc.add_paragraph()
            continue

        # Normal paragraph — handle inline bold
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        parts = re.split(r"(\*\*[^*]+\*\*)", stripped)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            else:
                r = p.add_run(part)
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)

# Process all 3 parts
for i, filepath in enumerate(FILES):
    print(f"Processing part {i+1}: {filepath}")
    parse_and_add(filepath)
    if i < len(FILES) - 1:
        add_page_break(doc)

doc.save(OUTPUT)
print(f"\n✅ Word document saved: {OUTPUT}")
